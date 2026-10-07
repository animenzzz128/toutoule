"""The export command (Task 1.12, step 5). No network and no real tracker."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_export_xlsx import TEMPLATE, add_job

from toutoule import config
from toutoule.cli import main
from toutoule.db import get_engine, get_session_factory, init_db


@pytest.fixture(autouse=True)
def no_real_profile(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The active profile is "sample" whatever is on the machine running the tests."""
    monkeypatch.setattr(config, "PROFILE_DIR", tmp_path / "absent")


@pytest.fixture
def database(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A throwaway database the CLI will open through DATABASE_URL."""
    path = tmp_path / "cli.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{path}")
    engine = get_engine(f"sqlite:///{path}")
    init_db(engine)
    return path


@pytest.fixture
def base(tmp_path: Path) -> Path:
    copy = tmp_path / "tracker.xlsx"
    shutil.copy(TEMPLATE, copy)
    return copy


def seed(database: Path, **kwargs: object) -> None:
    with get_session_factory(get_engine(f"sqlite:///{database}"))() as session:
        add_job(session, **kwargs)  # type: ignore[arg-type]


@pytest.mark.usefixtures("valid_env", "database")
def test_export_writes_the_file_and_reports_what_it_did(
    database: Path, base: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    seed(database)
    out = tmp_path / "export.xlsx"

    assert main(["export", "--base", str(base), "--out", str(out), "--sample"]) == 0
    printed = capsys.readouterr().out
    assert "added 1, skipped 0 already in tracker" in printed
    assert "was not changed" in printed
    assert load_workbook(out).worksheets[0].max_row == 4


@pytest.mark.usefixtures("valid_env", "database")
def test_jobs_filter_selects_one_job(
    database: Path, base: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    seed(database, company="First Co", url="https://example.com/1")
    seed(database, company="Second Co", url="https://example.com/2")
    out = tmp_path / "export.xlsx"

    assert main(["export", "--base", str(base), "--out", str(out), "--jobs", "1"]) == 0
    assert "added 1," in capsys.readouterr().out


@pytest.mark.usefixtures("valid_env", "database")
def test_a_refused_export_exits_one_and_writes_nothing(
    database: Path, base: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    seed(database, decisions=("rejected",))
    out = tmp_path / "export.xlsx"

    assert main(["export", "--base", str(base), "--out", str(out), "--jobs", "1"]) == 1
    assert "Export refused" in capsys.readouterr().err
    assert not out.exists()


@pytest.mark.usefixtures("valid_env", "database")
def test_a_bad_jobs_list_is_rejected_before_anything_opens(
    base: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "export.xlsx"
    assert main(["export", "--base", str(base), "--out", str(out), "--jobs", "1,two"]) == 1
    assert "comma-separated job ids" in capsys.readouterr().err


@pytest.mark.usefixtures("valid_env", "database")
def test_base_is_required(base: Path) -> None:
    # No default: a default path could quietly point at the owner's real tracker.
    with pytest.raises(SystemExit):
        main(["export"])


@pytest.mark.usefixtures("valid_env", "database")
def test_the_default_output_goes_under_data_private(
    database: Path, base: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    seed(database)
    monkeypatch.chdir(tmp_path)

    assert main(["export", "--base", str(base)]) == 0
    written = list((tmp_path / "data" / "private" / "exports").glob("tracker_*.xlsx"))
    assert len(written) == 1


@pytest.mark.usefixtures("valid_env", "database")
def test_the_output_names_each_job_it_added_and_skipped(
    database: Path, base: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    # An id on the command line says nothing about which posting it is, so the run names
    # what it acted on.
    seed(database, company="Northwind Labs", title="AI Product Manager")
    seed(database, company="Already There", title="Staff PM", url="https://example.com/jobs/1")
    out = tmp_path / "export.xlsx"

    assert main(["export", "--base", str(base), "--out", str(out)]) == 0
    printed = capsys.readouterr().out
    assert "added 1, skipped 1 already in tracker" in printed
    assert "  + Northwind Labs - AI Product Manager" in printed
    assert "  - Already There - Staff PM (already in tracker)" in printed
