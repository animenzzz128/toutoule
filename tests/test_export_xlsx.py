"""Writing the tracker file (Task 1.12, step 3). No network, no real tracker.

Every test works on a copy of tests/fixtures/tracker_template.xlsx, which carries the
owner's real header row, styling, dropdowns and conditional formatting.
"""

import hashlib
import json
import shutil
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from postings import extraction_answer
from sqlalchemy.orm import Session

from toutoule import config, export_xlsx, models
from toutoule.db import get_engine, get_session_factory, init_db
from toutoule.export_xlsx import COLUMNS, ExportRefused

TEMPLATE = Path(__file__).parent / "fixtures" / "tracker_template.xlsx"
SEEN_AT = datetime(2026, 10, 3, 15, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def no_real_profile(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Point the real-resume folder at a path that does not exist, in every test here.

    Without this, a machine that has data/private/profile would make the active profile
    "real" and these tests would depend on whose laptop they run on.
    """
    monkeypatch.setattr(config, "PROFILE_DIR", tmp_path / "absent")


@pytest.fixture
def session(tmp_path: Path) -> Iterator[Session]:
    """An empty database in a throwaway file."""
    engine = get_engine(f"sqlite:///{tmp_path / 'export.db'}")
    init_db(engine)
    with get_session_factory(engine)() as session:
        yield session


@pytest.fixture
def base(tmp_path: Path) -> Path:
    """A private copy of the committed template, so no test can alter the fixture."""
    copy = tmp_path / "tracker.xlsx"
    shutil.copy(TEMPLATE, copy)
    return copy


@pytest.fixture
def out(tmp_path: Path) -> Path:
    return tmp_path / "out" / "tracker_export.xlsx"


def match_payload(score: int = 72, version: str = "ai_product") -> dict[str, object]:
    """One scores.payload_json, shaped as score.save_scores writes it."""
    pair = {"requirement": "SQL required", "experience": "Python, SQL", "verified": True}
    return {
        "profile": "sample",
        "result": {
            "resume_version": version,
            "score": score,
            "dimensions": {
                "domain_fit": {"score": 8, "reason": "same work"},
                "skills_overlap": {"score": 7, "reason": "tools match"},
                "seniority_fit": {"score": 7, "reason": "right level"},
            },
            "evidence_pairs": [pair, pair, pair],
            "gaps": ["no people management", "no payments experience"],
            "unverified_pairs": 0,
            "recommended_version": version,
            "prompt_version": "score_v1",
            "model": "claude-sonnet-5-5",
        },
    }


def add_job(
    session: Session,
    *,
    company: str = "Example Co",
    title: str = "AI Product Manager",
    url: str = "https://example.com/job",
    score: int = 72,
    version: str = "ai_product",
    decisions: tuple[str, ...] = ("approved",),
    flags: tuple[tuple[str, str], ...] = (),
    extraction: bool = True,
    scored: bool = True,
    seen_at: datetime = SEEN_AT,
) -> models.Job:
    """One job with whatever an export needs, written straight to the database."""
    source = session.scalars(models.Source.__table__.select()).first()
    if source is None:
        session.add(models.Source(name="manual", tier=3, url="", adapter="manual"))
        session.commit()
    job = models.Job(
        source_id=1,
        company=company,
        title=title,
        url=url,
        raw_text="...",
        content_hash=f"hash-{company}-{title}",
        last_seen_at=seen_at,
    )
    session.add(job)
    session.commit()

    if extraction:
        session.add(
            models.Extraction(
                job_id=job.id,
                prompt_version="extract_v4",
                schema_version="1.0",
                payload_json=json.loads(extraction_answer()),
                model="claude-haiku-4-5",
                input_tokens=1000,
                output_tokens=300,
            )
        )
    if scored:
        session.add(
            models.Score(
                job_id=job.id,
                resume_version=version,
                score=score,
                payload_json=match_payload(score, version),
            )
        )
    for rule_id, severity in flags:
        session.add(
            models.RedFlag(job_id=job.id, rule_id=rule_id, severity=severity, evidence="quoted")
        )
    for action in decisions:
        session.add(
            models.Decision(
                job_id=job.id,
                action=action,
                reject_reason="not_interested" if action == "rejected" else None,
            )
        )
    session.commit()
    return job


def sheet(path: Path) -> Worksheet:
    return load_workbook(path).worksheets[0]


def value(worksheet: Worksheet, row: int, header: str) -> object:
    return worksheet.cell(row=row, column=COLUMNS.index(header) + 1).value


# --- the acceptance criterion -------------------------------------------------------------


def test_headers_survive_the_round_trip(session: Session, base: Path, out: Path) -> None:
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    assert [cell.value for cell in sheet(out)[1]] == list(COLUMNS)


def test_formatting_is_preserved(session: Session, base: Path, out: Path) -> None:
    add_job(session)
    before, _ = sheet(base), export_xlsx.export_jobs(session, base, out)
    after = sheet(out)

    assert after.freeze_panes == before.freeze_panes == "C2"
    assert {k: d.width for k, d in after.column_dimensions.items()} == {
        k: d.width for k, d in before.column_dimensions.items()
    }
    assert [str(d.sqref) for d in after.data_validations.dataValidation] == [
        str(d.sqref) for d in before.data_validations.dataValidation
    ]
    assert [str(r.sqref) for r in after.conditional_formatting] == [
        str(r.sqref) for r in before.conditional_formatting
    ]
    assert after.tables["JobApplicationTracker"].tableStyleInfo.name == "TableStyleMedium2"
    assert [cell.font.b for cell in after[1]] == [True] * 18


# --- what lands in the cells --------------------------------------------------------------


def test_values_land_under_the_right_headers(session: Session, base: Path, out: Path) -> None:
    add_job(session, score=84)
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    assert new.max_row == 4  # the template's two fake rows, plus ours

    assert value(new, 4, "Company") == "Example Co"
    assert value(new, 4, "Title") == "AI Product Manager"
    assert value(new, 4, "Status") == "To Apply"
    assert value(new, 4, "Link") == "https://example.com/job"
    assert value(new, 4, "Geography") == "New York, NY"
    assert value(new, 4, "Deadline") == "2026-12-31"
    assert value(new, 4, "Candidate Fit") == 84
    assert value(new, 4, "Priority Score") is None
    assert value(new, 4, "Main Skills Required") == "SQL"
    assert value(new, 4, "Recommended Resume") == "ai_product"
    assert value(new, 4, "Notes") is None
    for header in ("Market", "Application Date", "Interview Stage", "Interviewer", "Next Action"):
        assert value(new, 4, header) is None


def test_exported_status_is_one_of_the_columns_own_dropdown_values(
    session: Session, base: Path, out: Path
) -> None:
    # Read the allowed list out of the file itself, so the assertion cannot drift from it.
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    status_column = f"{chr(ord('A') + COLUMNS.index('Status'))}2"
    allowed = next(
        d.formula1.strip('"').split(",")
        for d in new.data_validations.dataValidation
        if status_column in str(d.sqref)
    )
    assert value(new, 4, "Status") in allowed


def test_last_verified_is_a_date_with_the_trackers_date_format(
    session: Session, base: Path, out: Path
) -> None:
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    cell = new.cell(row=4, column=COLUMNS.index("Last Verified") + 1)
    assert cell.value == datetime(2026, 10, 3)  # 15:00 UTC is 11:00 in New York
    assert cell.number_format == new.cell(row=2, column=7).number_format


def test_notes_record_the_flags(session: Session, base: Path, out: Path) -> None:
    add_job(session, flags=(("R1", "HARD"), ("R5", "URGENT")))
    export_xlsx.export_jobs(session, base, out)
    assert value(sheet(out), 4, "Notes") == (
        "toutoule flags: HARD R1; URGENT R5 (approved after HARD confirm)"
    )


# --- styling ------------------------------------------------------------------------------


def test_new_cells_have_no_fill(session: Session, base: Path, out: Path) -> None:
    # The table's banding and the conditional formatting own every colour in the tracker.
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    assert [cell.fill.fill_type for cell in new[4]] == [None] * 18


def test_new_row_copies_font_alignment_and_height(session: Session, base: Path, out: Path) -> None:
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    for column in range(1, 19):
        written, model = new.cell(row=4, column=column), new.cell(row=2, column=column)
        assert written.font.name == model.font.name
        assert written.alignment.wrap_text == model.alignment.wrap_text
    assert new.row_dimensions[4].height == new.row_dimensions[2].height


def test_table_ref_covers_every_exported_row(session: Session, base: Path, out: Path) -> None:
    add_job(session)
    add_job(session, company="Second Co", url="https://example.com/two")
    export_xlsx.export_jobs(session, base, out)
    new = sheet(out)
    table = new.tables["JobApplicationTracker"]
    assert table.ref == "A1:R5"
    assert table.autoFilter.ref == "A1:R5"


# --- refusals -----------------------------------------------------------------------------


def test_refuses_to_write_over_the_file_it_reads(session: Session, base: Path) -> None:
    add_job(session)
    with pytest.raises(ExportRefused, match="the one it reads"):
        export_xlsx.export_jobs(session, base, base)


def test_the_base_file_is_never_modified(session: Session, base: Path, out: Path) -> None:
    add_job(session)
    before = hashlib.sha256(base.read_bytes()).hexdigest()
    export_xlsx.export_jobs(session, base, out)
    assert hashlib.sha256(base.read_bytes()).hexdigest() == before


def test_a_renamed_header_is_refused_and_nothing_is_written(
    session: Session, base: Path, out: Path
) -> None:
    workbook = load_workbook(base)
    workbook.worksheets[0]["F1"] = "Location"
    workbook.save(base)
    add_job(session)

    with pytest.raises(ExportRefused, match="column 6 .* 'Location'.*'Geography'"):
        export_xlsx.export_jobs(session, base, out)
    assert not out.exists()


def test_a_nineteenth_column_is_refused(session: Session, base: Path, out: Path) -> None:
    workbook = load_workbook(base)
    workbook.worksheets[0]["S1"] = "Salary"
    workbook.save(base)
    add_job(session)

    with pytest.raises(ExportRefused, match="19th column"):
        export_xlsx.export_jobs(session, base, out)


def test_a_named_job_that_is_not_approved_is_refused(
    session: Session, base: Path, out: Path
) -> None:
    job = add_job(session, decisions=("rejected",))
    with pytest.raises(ExportRefused, match="no approved decision"):
        export_xlsx.export_jobs(session, base, out, job_ids=[job.id])
    assert not out.exists()


def test_a_job_without_a_score_is_refused(session: Session, base: Path, out: Path) -> None:
    add_job(session, scored=False)
    with pytest.raises(ExportRefused, match="has no sample score"):
        export_xlsx.export_jobs(session, base, out)
    assert not out.exists()


def test_a_job_without_an_extraction_is_refused(session: Session, base: Path, out: Path) -> None:
    add_job(session, extraction=False)
    with pytest.raises(ExportRefused, match="no extraction"):
        export_xlsx.export_jobs(session, base, out)


def test_a_locked_output_file_says_how_to_fix_it(
    session: Session, base: Path, out: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def locked(self: Workbook, path: object) -> None:
        raise PermissionError(13, "Permission denied")

    add_job(session)
    monkeypatch.setattr(Workbook, "save", locked)
    with pytest.raises(ExportRefused, match="Close .* in Excel and try again"):
        export_xlsx.export_jobs(session, base, out)


# --- selection ----------------------------------------------------------------------------


def test_only_approved_jobs_are_exported(session: Session, base: Path, out: Path) -> None:
    add_job(session, company="Yes Co", url="https://example.com/yes")
    add_job(session, company="No Co", url="https://example.com/no", decisions=("rejected",))
    add_job(session, company="Undecided Co", url="https://example.com/un", decisions=())

    summary = export_xlsx.export_jobs(session, base, out)
    assert summary.added == 1
    assert value(sheet(out), 4, "Company") == "Yes Co"


def test_the_latest_decision_is_the_one_that_counts(
    session: Session, base: Path, out: Path
) -> None:
    # decisions is append-only: approving then rejecting leaves both rows, and the job
    # must not be exported.
    add_job(session, company="Changed Mind", decisions=("approved", "rejected"))
    assert export_xlsx.export_jobs(session, base, out).added == 0


def test_a_rejection_followed_by_an_approval_is_exported(
    session: Session, base: Path, out: Path
) -> None:
    add_job(session, company="Reconsidered", decisions=("rejected", "approved"))
    assert export_xlsx.export_jobs(session, base, out).added == 1


# --- not writing the same job twice -------------------------------------------------------


def test_a_job_whose_link_is_already_in_the_tracker_is_skipped(
    session: Session, base: Path, out: Path
) -> None:
    add_job(session, company="Anything", url="https://example.com/jobs/1")
    summary = export_xlsx.export_jobs(session, base, out)
    assert (summary.added, summary.skipped) == (0, 1)
    assert sheet(out).max_row == 3


def test_company_and_title_decide_when_a_link_is_missing(
    session: Session, base: Path, out: Path
) -> None:
    # The template's row 3 is Sample Tech / AI Product Manager with a link; blank the link
    # so the comparison has to fall back to the name, ignoring case and spacing.
    workbook = load_workbook(base)
    workbook.worksheets[0]["D3"] = None
    workbook.save(base)
    add_job(session, company="  SAMPLE TECH ", title="ai product manager", url="")

    assert export_xlsx.export_jobs(session, base, out).skipped == 1


def test_exporting_the_same_job_into_an_output_that_has_it_adds_nothing(
    session: Session, base: Path, out: Path, tmp_path: Path
) -> None:
    add_job(session)
    first = export_xlsx.export_jobs(session, base, out)
    assert first.added == 1

    second = export_xlsx.export_jobs(session, out, tmp_path / "again.xlsx")
    assert (second.added, second.skipped) == (0, 1)
    assert sheet(tmp_path / "again.xlsx").max_row == 4


# --- keeping the dropdowns and colour rules over the new rows -----------------------------


def fill_to_row(base: Path, last_row: int) -> None:
    """Pad the base file with filler rows so the next export crosses row 200."""
    workbook = load_workbook(base)
    worksheet = workbook.worksheets[0]
    for row in range(4, last_row + 1):
        worksheet.cell(row=row, column=1, value=f"Filler {row}")
        worksheet.cell(row=row, column=2, value="Role")
        worksheet.cell(row=row, column=4, value=f"https://example.com/filler/{row}")
    worksheet.tables["JobApplicationTracker"].ref = f"A1:R{last_row}"
    workbook.save(base)


def ranges(worksheet: Worksheet) -> tuple[list[str], list[str]]:
    return (
        [str(d.sqref) for d in worksheet.data_validations.dataValidation],
        [str(r.sqref) for r in worksheet.conditional_formatting],
    )


def test_ranges_are_left_alone_when_the_new_rows_already_fit(
    session: Session, base: Path, out: Path
) -> None:
    before = ranges(sheet(base))
    add_job(session)
    export_xlsx.export_jobs(session, base, out)
    assert ranges(sheet(out)) == before
    assert ranges(sheet(out))[0] == ["C2:C200", "E2:E200", "H2:H200"]


def test_ranges_grow_once_an_export_passes_the_row_they_stop_at(
    session: Session, base: Path, out: Path
) -> None:
    fill_to_row(base, 200)
    add_job(session, company="Row 201 Co", url="https://example.com/201")
    export_xlsx.export_jobs(session, base, out)

    validations, formats = ranges(sheet(out))
    assert validations == ["C2:C201", "E2:E201", "H2:H201"]
    assert formats == ["C2:C201", "N2:N201"]
    assert sheet(out).tables["JobApplicationTracker"].ref == "A1:R201"


def test_a_grown_dropdown_keeps_its_allowed_values(session: Session, base: Path, out: Path) -> None:
    fill_to_row(base, 200)
    add_job(session, company="Row 201 Co", url="https://example.com/201")
    export_xlsx.export_jobs(session, base, out)

    status = next(
        d for d in sheet(out).data_validations.dataValidation if str(d.sqref).startswith("C2")
    )
    assert "To Apply" in status.formula1
    assert value(sheet(out), 201, "Status") == "To Apply"


def test_a_grown_colour_rule_keeps_its_rules(session: Session, base: Path, out: Path) -> None:
    fill_to_row(base, 200)
    add_job(session, company="Row 201 Co", url="https://example.com/201")
    export_xlsx.export_jobs(session, base, out)

    by_range = {str(entry.sqref): entry for entry in sheet(out).conditional_formatting}
    assert [rule.type for rule in by_range["C2:C201"].rules] == ["expression"] * 3
    assert [rule.type for rule in by_range["N2:N201"].rules] == ["colorScale"]
    assert by_range["N2:N201"].rules[0].colorScale is not None
