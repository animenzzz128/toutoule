"""Export approved jobs into the owner's existing 18-column tracker (PRD F8, D-011).

The tracker is the owner's file and the system is a guest in it. Six columns are his
alone and are always written blank; one more, Priority Score, stays blank because the
match score is a measure of fit and must never be presented as a priority (D-010).

The first half of the module is pure: it turns stored rows into the eighteen values a
spreadsheet row holds, and knows nothing about openpyxl, files or styling. The second
half opens the owner's file, appends to it, and never writes over the file it read.
"""

import logging
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import NamedTuple
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from toutoule import models, schemas
from toutoule.score import ProfileKind
from toutoule.triage import active_profile_dir

logger = logging.getLogger(__name__)

# The owner's header row, in his order. The export never invents, reorders or renames a
# column; it checks the file it is given against this and refuses if they differ.
COLUMNS: tuple[str, ...] = (
    "Company",
    "Title",
    "Status",
    "Link",
    "Market",
    "Geography",
    "Application Date",
    "Interview Stage",
    "Interviewer",
    "Notes",
    "Deadline",
    "Urgency",
    "Candidate Fit",
    "Priority Score",
    "Main Skills Required",
    "Recommended Resume",
    "Next Action",
    "Last Verified",
)

# What a field says when the posting did not state it. Never a guess, never a blank cell:
# a blank is ambiguous between "the posting was silent" and "the system never looked"
# (PRD F3, non-negotiable constraint 2).
NOT_STATED = "Not stated"

# The one value the export writes into Status. The tracker's vocabulary is the owner's,
# not JobStatus: "approved but not yet applied" is what he calls "To Apply". It is also a
# member of the column's dropdown, which a test checks against the file itself.
TO_APPLY = "To Apply"

# jobs.last_seen_at is UTC. The tracker is read by a human in New York, so a job triaged
# at 21:00 on the 5th must not be stamped the 6th (the DST hazard in CLAUDE.md, in small).
TRACKER_TZ = ZoneInfo("America/New_York")


def field_text(field: schemas.ExtractedField) -> str:
    """One extracted field as a cell value: its verbatim value, or "Not stated".

    Written against ExtractedField rather than against any one field name, so the rule is
    identical for all five critical and all five important fields.
    """
    if not field.stated or field.value is None:
        return NOT_STATED
    return field.value


def notes_line(flags: list[models.RedFlag]) -> str:
    """One line naming every rule that fired, or "" when none did.

    Only approved jobs are exported, and approving past a HARD flag already requires an
    explicit confirmation (triage.record_decision), so a HARD flag on an exported job is
    by construction one the owner overrode on purpose. The suffix records that in the
    tracker, where he will read it months later.
    """
    if not flags:
        return ""
    ordered = sorted(flags, key=lambda flag: flag.rule_id)
    body = "; ".join(f"{flag.severity} {flag.rule_id}" for flag in ordered)
    line = f"toutoule flags: {body}"
    if any(flag.severity == models.FlagSeverity.HARD for flag in ordered):
        line += " (approved after HARD confirm)"
    return line


def last_verified(seen_at: datetime) -> date:
    """The tracker's "Last Verified" date: when the system last saw the posting, in New York.

    ZoneInfo is the standard library's time-zone database; .astimezone moves the instant
    into that zone before the date is taken, which is the whole point of the conversion.
    """
    return seen_at.astimezone(TRACKER_TZ).date()


def build_row(
    job: models.Job,
    extraction: schemas.Extraction,
    result: schemas.MatchResult,
    flags: list[models.RedFlag],
) -> list[object]:
    """The eighteen values for one job, in COLUMNS order. None means an empty cell.

    Six columns are the owner's to fill (Application Date, Interview Stage, Interviewer,
    Next Action, Urgency, Market) and one is blank by decision (Priority Score, D-010).
    Market has no source at all: the market is an argument to the red-flag rules and is
    never persisted, so there is nothing to write (D-011, README known limits).
    """
    return [
        job.company,  # Company
        job.title,  # Title
        TO_APPLY,  # Status
        job.url,  # Link
        None,  # Market — not persisted
        field_text(extraction.important.location),  # Geography
        None,  # Application Date — owner only
        None,  # Interview Stage — owner only
        None,  # Interviewer — owner only
        notes_line(flags) or None,  # Notes
        field_text(extraction.critical.deadline),  # Deadline — verbatim, never parsed
        None,  # Urgency — owner only; the flags are in Notes
        result.score,  # Candidate Fit — the match score belongs here, not in Priority
        None,  # Priority Score — blank by D-010
        ", ".join(extraction.reference.skills) or None,  # Main Skills Required
        result.recommended_version,  # Recommended Resume
        None,  # Next Action — owner only
        last_verified(job.last_seen_at),  # Last Verified
    ]


# --- Reading what the database already decided --------------------------------------------


class ExportRefused(Exception):
    """The export would be wrong or unsafe. Nothing was written; the message says why."""


class ExportSummary(NamedTuple):
    """What one export did. Printed by the CLI and asserted by the tests."""

    added: int
    skipped: int
    path: Path


def approved_job_ids(session: Session) -> list[int]:
    """Jobs whose most recent decision is "approved".

    decisions is append-only, so a job the owner approved and later rejected has two rows
    and only the newest one counts. The subquery picks the highest id per job, which is
    the newest because ids ascend with insertion.
    """
    newest = (
        select(models.Decision.job_id, func.max(models.Decision.id).label("last_id"))
        .group_by(models.Decision.job_id)
        .subquery()
    )
    statement = (
        select(models.Decision.job_id)
        .join(newest, models.Decision.id == newest.c.last_id)
        .where(models.Decision.action == "approved")
        .order_by(models.Decision.job_id)
    )
    return list(session.scalars(statement))


def _latest_extraction(session: Session, job_id: int) -> schemas.Extraction | None:
    row = session.scalars(
        select(models.Extraction)
        .where(models.Extraction.job_id == job_id)
        .order_by(models.Extraction.id.desc())
    ).first()
    return schemas.Extraction.model_validate(row.payload_json) if row else None


def _recommended_result(
    session: Session, job_id: int, profile: ProfileKind
) -> schemas.MatchResult | None:
    """The winning version's score, from this profile's rows only.

    Scores from the real resumes and from the samples are different numbers about
    different documents, so one is never shown under the other's name (as in triage.py).
    """
    rows = session.scalars(
        select(models.Score).where(models.Score.job_id == job_id).order_by(models.Score.id)
    ).all()
    newest = {row.resume_version: row for row in rows if row.payload_json.get("profile") == profile}
    for row in newest.values():
        result = schemas.MatchResult.model_validate(row.payload_json["result"])
        if result.resume_version == result.recommended_version:
            return result
    return None


def _flags(session: Session, job_id: int) -> list[models.RedFlag]:
    return list(session.scalars(select(models.RedFlag).where(models.RedFlag.job_id == job_id)))


# --- Guards on the file we were handed ----------------------------------------------------


def check_headers(worksheet: Worksheet) -> None:
    """Refuse any file whose header row is not the owner's, naming the first difference.

    The export appends into a layout it did not create. If the columns have moved, every
    value afterwards lands in the wrong place, and a spreadsheet full of plausible wrong
    answers is worse than no spreadsheet.
    """
    for index, expected in enumerate(COLUMNS, start=1):
        found = worksheet.cell(row=1, column=index).value
        if found != expected:
            raise ExportRefused(
                f"column {index} of {worksheet.title!r} is {found!r}, expected {expected!r}; "
                "the export writes the owner's layout or nothing"
            )
    extra = worksheet.cell(row=1, column=len(COLUMNS) + 1).value
    if extra is not None:
        raise ExportRefused(f"the tracker has a 19th column, {extra!r}; expected exactly 18")


class _Existing(NamedTuple):
    """The keys already in the tracker, used to avoid writing a job in twice."""

    urls: set[str]
    names_without_url: set[tuple[str, str]]
    all_names: set[tuple[str, str]]


def _name_key(company: object, title: object) -> tuple[str, str]:
    return (str(company or "").strip().lower(), str(title or "").strip().lower())


def _existing_keys(worksheet: Worksheet) -> _Existing:
    urls: set[str] = set()
    names_without_url: set[tuple[str, str]] = set()
    all_names: set[tuple[str, str]] = set()
    for company, title, _status, link, *_rest in worksheet.iter_rows(min_row=2, values_only=True):
        if company is None and title is None and link is None:
            continue  # a styled but empty row
        name = _name_key(company, title)
        all_names.add(name)
        url = str(link or "").strip()
        if url:
            urls.add(url)
        else:
            names_without_url.add(name)
    return _Existing(urls, names_without_url, all_names)


def _is_duplicate(job: models.Job, existing: _Existing) -> bool:
    """True if this job is already in the tracker.

    The link is the identity when both sides have one. When either side does not, company
    and title decide instead, trimmed and compared without case.
    """
    url = (job.url or "").strip()
    name = _name_key(job.company, job.title)
    if url:
        return url in existing.urls or name in existing.names_without_url
    return name in existing.all_names


# --- Writing ------------------------------------------------------------------------------


def _copy_style(source: Cell, target: Cell) -> None:
    """Give the new cell the base row's look — but never its fill.

    The table's banding and the conditional formatting supply every colour the tracker is
    meant to have. Copying a fill would make a one-off highlight in the base row permanent
    on every future row. copy() is needed because openpyxl style objects are shared: an
    uncopied assignment would leave many cells pointing at one record.
    """
    target.font = copy(source.font)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format


def _with_last_row(ref: str, last_row: int) -> str:
    """The same rectangle, ending at last_row. "A1:R3" with 7 becomes "A1:R7"."""
    min_col, min_row, max_col, _ = range_boundaries(ref)
    assert min_col and min_row and max_col  # a parsed ref always has these
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{last_row}"


def _widen_table(worksheet: Worksheet, last_row: int) -> None:
    """Grow the table (and any filter) to cover the rows just appended.

    openpyxl does not do this: a table whose ref stops short leaves the new rows outside
    it, without banding and outside the filter.
    """
    for table in worksheet.tables.values():
        table.ref = _with_last_row(table.ref, last_row)
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    if worksheet.auto_filter.ref:  # a base file that still has a sheet-level filter
        worksheet.auto_filter.ref = _with_last_row(worksheet.auto_filter.ref, last_row)


def _extended_sqref(ranges: MultiCellRange, last_row: int) -> str | None:
    """The same ranges with any that stop short grown to last_row, or None if none did.

    Returning None for "nothing to do" is what keeps an export that fits inside the
    existing ranges from rewriting them at all.
    """
    parts: list[str] = []
    grew = False
    for cell_range in ranges:
        max_row = cell_range.max_row
        if max_row < last_row:
            max_row, grew = last_row, True
        parts.append(
            f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}:"
            f"{get_column_letter(cell_range.max_col)}{max_row}"
        )
    return " ".join(parts) if grew else None


def _extend_validations(worksheet: Worksheet, last_row: int) -> None:
    """Keep the dropdowns covering every row. They stop at row 200 in the owner's file."""
    for validation in worksheet.data_validations.dataValidation:
        grown = _extended_sqref(validation.sqref, last_row)
        if grown is not None:
            validation.sqref = MultiCellRange(grown)


def _extend_conditional_formatting(worksheet: Worksheet, last_row: int) -> None:
    """Keep the colour rules covering every row, by rebuilding the collection.

    openpyxl keys conditional formatting by its range, so a range cannot simply be edited
    in place; the rules are re-added under the grown range instead. The whole collection
    is left alone unless at least one range stopped short.
    """
    entries = list(worksheet.conditional_formatting)
    if not any(_extended_sqref(entry.sqref, last_row) for entry in entries):
        return
    rebuilt = ConditionalFormattingList()
    for entry in entries:
        grown = _extended_sqref(entry.sqref, last_row) or str(entry.sqref)
        for rule in entry.rules:
            rebuilt.add(grown, rule)
    worksheet.conditional_formatting = rebuilt


def export_jobs(
    session: Session,
    base_path: Path,
    out_path: Path,
    job_ids: list[int] | None = None,
    *,
    sample: bool = False,
) -> ExportSummary:
    """Append every approved job to a copy of the owner's tracker. Returns what it did.

    base_path is read and never written. out_path is a new file, so a failed export cannot
    damage the tracker and the owner decides when to replace it.

    Refuses, before writing anything, if: the two paths are the same file; a named job is
    not approved; an approved job has no extraction or no score for the active profile; or
    the base file's header row is not the owner's eighteen columns.
    """
    base_path, out_path = Path(base_path), Path(out_path)
    if base_path.resolve() == out_path.resolve():
        raise ExportRefused(f"the export writes a new file; {base_path} is the one it reads")

    approved = approved_job_ids(session)
    if job_ids is None:
        wanted = approved
    else:
        not_approved = [job_id for job_id in job_ids if job_id not in approved]
        if not_approved:
            raise ExportRefused(
                f"jobs {not_approved} have no approved decision; only approved jobs are exported"
            )
        wanted = list(job_ids)

    profile = active_profile_dir(sample).kind
    workbook = load_workbook(base_path)
    worksheet = workbook.worksheets[0]
    check_headers(worksheet)

    # Every row is built first, so a job missing its extraction or score stops the export
    # before a partial file exists.
    prepared: list[tuple[models.Job, list[object]]] = []
    for job_id in wanted:
        job = session.get(models.Job, job_id)
        if job is None:
            raise ExportRefused(f"no job with id {job_id}")
        extraction = _latest_extraction(session, job_id)
        if extraction is None:
            raise ExportRefused(f"job {job_id} ({job.company}) has no extraction to export")
        result = _recommended_result(session, job_id, profile)
        if result is None:
            raise ExportRefused(
                f"job {job_id} ({job.company}) has no {profile} score; "
                "score it before exporting, or the fit column would be a blank guess"
            )
        prepared.append((job, build_row(job, extraction, result, _flags(session, job_id))))

    existing = _existing_keys(worksheet)
    style_row = 2 if worksheet.max_row >= 2 else None
    date_format = _date_format(worksheet, style_row)
    row_number = worksheet.max_row
    added = skipped = 0

    for job, values in prepared:
        if _is_duplicate(job, existing):
            logger.info("job %s (%s) is already in the tracker", job.id, job.company)
            skipped += 1
            continue
        row_number += 1
        _write_row(worksheet, row_number, values, style_row, date_format)
        existing.urls.add((job.url or "").strip())
        existing.all_names.add(_name_key(job.company, job.title))
        added += 1

    if added:
        _widen_table(worksheet, row_number)
        _extend_validations(worksheet, row_number)
        _extend_conditional_formatting(worksheet, row_number)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(out_path)
    except PermissionError as error:
        raise ExportRefused(f"Close {out_path} in Excel and try again.") from error
    logger.info("added %s, skipped %s already in tracker", added, skipped)
    return ExportSummary(added=added, skipped=skipped, path=out_path)


def _date_format(worksheet: Worksheet, style_row: int | None) -> str:
    """The tracker's own date format, read off Application Date, for Last Verified.

    Taking it from the file rather than hard-coding one means the exported date matches
    whatever the owner's other date column already uses.
    """
    if style_row is None:
        return "yyyy-mm-dd"
    found = worksheet.cell(row=style_row, column=COLUMNS.index("Application Date") + 1)
    return "yyyy-mm-dd" if found.number_format == "General" else found.number_format


def _write_row(
    worksheet: Worksheet,
    row_number: int,
    values: list[object],
    style_row: int | None,
    date_format: str,
) -> None:
    """Write one row's values and give it the base row's look and height."""
    for column_index, value in enumerate(values, start=1):
        cell = worksheet.cell(row=row_number, column=column_index, value=value)
        if style_row is not None:
            _copy_style(worksheet.cell(row=style_row, column=column_index), cell)
    # Last Verified holds a real date, so it needs a date format even where the base row's
    # own cell did not have one.
    worksheet.cell(
        row=row_number, column=COLUMNS.index("Last Verified") + 1
    ).number_format = date_format
    if style_row is not None:
        worksheet.row_dimensions[row_number].height = worksheet.row_dimensions[style_row].height
